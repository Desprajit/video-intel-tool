from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
import datetime


# ── Color Palette ──────────────────────────────────────────────
BG_DARK      = RGBColor(0x0D, 0x0F, 0x1A)   # slide background
BG_CARD      = RGBColor(0x17, 0x1B, 0x2E)   # card / section bg
ACCENT       = RGBColor(0x4F, 0x8E, 0xFF)   # primary blue
ACCENT2      = RGBColor(0x00, 0xE5, 0xB8)   # teal highlight
ACCENT3      = RGBColor(0xFF, 0x6B, 0x6B)   # red/coral
ACCENT4      = RGBColor(0xFF, 0xC1, 0x5E)   # amber
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
GRAY         = RGBColor(0x8A, 0x8F, 0xA8)
LIGHT_GRAY   = RGBColor(0x1E, 0x24, 0x3B)

COMPANY_COLORS = [
    ACCENT,
    ACCENT2,
    ACCENT3,
    ACCENT4,
    RGBColor(0xA0, 0x6C, 0xFF)
]

W = Inches(13.33)
H = Inches(7.5)


def hex_to_rgb(hex_str):
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def safe_str(value, default=""):
    if value is None:
        return default
    return str(value)


def fmt_num(n):
    try:
        n = float(n)
    except (TypeError, ValueError):
        return "0"

    if n >= 1_000_000:
        return f"{n / 1_000_000:.1f}M"
    elif n >= 1_000:
        return f"{n / 1_000:.1f}K"
    return str(int(n))


# ── Low-level helpers ───────────────────────────────────────────

def add_rect(slide, x, y, w, h, fill=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.line.width = line_width

    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()

    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width if line_width else Pt(1)
    else:
        shape.line.fill.background()

    return shape


def add_text(slide, text, x, y, w, h, size=Pt(12), bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(x, y, w, h)
    tf = txb.text_frame
    tf.word_wrap = wrap

    p = tf.paragraphs[0]
    p.alignment = align

    run = p.add_run()
    run.text = safe_str(text)
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color

    return txb


def set_slide_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_slide(prs, layout_idx=6):
    layout = prs.slide_layouts[layout_idx]
    slide = prs.slides.add_slide(layout)
    set_slide_bg(slide)

    for ph in slide.placeholders:
        sp = ph._element
        sp.getparent().remove(sp)

    return slide


def add_accent_bar(slide, color=ACCENT):
    add_rect(slide, 0, Inches(7.2), W, Inches(0.3), fill=color)


def section_header(slide, label, color=ACCENT):
    add_rect(slide, Inches(0.4), Inches(0.3), Inches(0.07), Inches(0.5), fill=color)
    add_text(
        slide,
        label.upper(),
        Inches(0.6),
        Inches(0.28),
        Inches(8),
        Inches(0.5),
        size=Pt(9),
        bold=True,
        color=color
    )


def slide_title(slide, title, subtitle=None):
    add_text(
        slide,
        title,
        Inches(0.5),
        Inches(0.7),
        Inches(12.3),
        Inches(0.9),
        size=Pt(28),
        bold=True,
        color=WHITE
    )
    if subtitle:
        add_text(
            slide,
            subtitle,
            Inches(0.5),
            Inches(1.45),
            Inches(12.3),
            Inches(0.45),
            size=Pt(13),
            color=GRAY
        )


def metric_card(slide, x, y, w, h, label, value, sub=None, color=ACCENT):
    add_rect(slide, x, y, w, h, fill=BG_CARD, line_color=color, line_width=Pt(1))
    add_text(
        slide,
        label,
        x + Inches(0.15),
        y + Inches(0.12),
        w - Inches(0.3),
        Inches(0.3),
        size=Pt(9),
        color=GRAY
    )
    add_text(
        slide,
        safe_str(value),
        x + Inches(0.15),
        y + Inches(0.35),
        w - Inches(0.3),
        Inches(0.5),
        size=Pt(22),
        bold=True,
        color=color
    )
    if sub:
        add_text(
            slide,
            sub,
            x + Inches(0.15),
            y + Inches(0.8),
            w - Inches(0.3),
            Inches(0.25),
            size=Pt(9),
            color=GRAY
        )


# ── Slides ───────────────────────────────────────────────────────

def slide_cover(prs, report_data):
    slide = add_slide(prs)
    add_rect(slide, 0, 0, W, H, fill=BG_DARK)
    add_rect(slide, 0, 0, Inches(5), H, fill=RGBColor(0x11, 0x16, 0x26))
    add_rect(slide, 0, H - Inches(0.5), W, Inches(0.5), fill=ACCENT)

    # decorative blocks
    add_rect(slide, Inches(9.2), Inches(0.3), Inches(3.4), Inches(2.2), fill=RGBColor(0x1A, 0x22, 0x40))
    add_rect(slide, Inches(10.4), Inches(4.0), Inches(2.8), Inches(2.1), fill=RGBColor(0x13, 0x1A, 0x33))

    add_text(
        slide,
        "VIDEO COMPETITOR",
        Inches(0.6),
        Inches(1.2),
        Inches(9),
        Inches(0.8),
        size=Pt(42),
        bold=True,
        color=WHITE
    )
    add_text(
        slide,
        "INTELLIGENCE REPORT",
        Inches(0.6),
        Inches(1.9),
        Inches(9),
        Inches(0.8),
        size=Pt(42),
        bold=True,
        color=ACCENT
    )

    companies = report_data.get("companies", [])
    main = report_data.get("main_company", "")
    company_names = " · ".join([c.get("company", "") for c in companies if "error" not in c])

    add_text(
        slide,
        company_names,
        Inches(0.6),
        Inches(2.9),
        Inches(10),
        Inches(0.5),
        size=Pt(14),
        color=GRAY
    )

    add_text(
        slide,
        f"Prepared for: {main}",
        Inches(0.6),
        Inches(3.5),
        Inches(8),
        Inches(0.4),
        size=Pt(12),
        color=ACCENT2
    )
    add_text(
        slide,
        f"Report Date: {datetime.date.today().strftime('%B %d, %Y')}",
        Inches(0.6),
        Inches(3.95),
        Inches(8),
        Inches(0.35),
        size=Pt(11),
        color=GRAY
    )

    add_text(
        slide,
        "Powered by YouTube Data API + Gemini AI",
        Inches(0.6),
        Inches(6.8),
        Inches(8),
        Inches(0.35),
        size=Pt(9),
        color=GRAY
    )


def slide_executive_summary(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Executive Summary", ACCENT2)
    slide_title(slide, "Who Is Leading in Video Marketing?")
    add_accent_bar(slide, ACCENT2)

    insights = report_data.get("insights", {})
    summary = insights.get("executive_summary", "Analysis complete.")
    leader = insights.get("leader", "")
    leader_reason = insights.get("leader_reason", "")

    add_rect(slide, Inches(0.4), Inches(1.9), Inches(8), Inches(2), fill=BG_CARD)
    add_text(
        slide,
        summary,
        Inches(0.6),
        Inches(2.0),
        Inches(7.7),
        Inches(1.8),
        size=Pt(13),
        color=WHITE,
        wrap=True
    )

    add_rect(
        slide,
        Inches(9),
        Inches(1.9),
        Inches(3.9),
        Inches(2),
        fill=ACCENT,
        line_color=ACCENT2,
        line_width=Pt(1.5)
    )
    add_text(
        slide,
        "🏆  LEADER",
        Inches(9.1),
        Inches(2.0),
        Inches(3.7),
        Inches(0.4),
        size=Pt(10),
        bold=True,
        color=ACCENT2
    )
    add_text(
        slide,
        leader,
        Inches(9.1),
        Inches(2.35),
        Inches(3.7),
        Inches(0.6),
        size=Pt(20),
        bold=True,
        color=WHITE
    )
    add_text(
        slide,
        leader_reason,
        Inches(9.1),
        Inches(2.85),
        Inches(3.6),
        Inches(0.8),
        size=Pt(10),
        color=WHITE,
        wrap=True
    )

    companies = [c for c in report_data.get("companies", []) if "error" not in c]
    x_positions = [Inches(0.4), Inches(3.35), Inches(6.3), Inches(9.25), Inches(12.2)]

    for i, comp in enumerate(companies[:5]):
        cx = x_positions[i]
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]

        add_rect(
            slide,
            cx,
            Inches(4.2),
            Inches(2.6),
            Inches(1.5),
            fill=BG_CARD,
            line_color=color,
            line_width=Pt(1)
        )
        add_text(
            slide,
            comp.get("company", "")[:18],
            cx + Inches(0.1),
            Inches(4.28),
            Inches(2.4),
            Inches(0.35),
            size=Pt(10),
            bold=True,
            color=color
        )
        add_text(
            slide,
            fmt_num(comp.get("subscribers", 0)) + " subs",
            cx + Inches(0.1),
            Inches(4.6),
            Inches(2.4),
            Inches(0.35),
            size=Pt(13),
            bold=True,
            color=WHITE
        )
        add_text(
            slide,
            f"Avg {fmt_num(comp.get('avg_views', 0))} views",
            cx + Inches(0.1),
            Inches(4.9),
            Inches(2.4),
            Inches(0.3),
            size=Pt(10),
            color=GRAY
        )


def slide_channel_overview(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Channel Overview", ACCENT)
    slide_title(slide, "Channel Comparison")
    add_accent_bar(slide)

    companies = [c for c in report_data.get("companies", []) if "error" not in c]
    headers = ["Company", "Channel", "Subscribers", "Total Videos", "Total Views", "Since"]
    col_w = [Inches(1.8), Inches(2.2), Inches(1.6), Inches(1.6), Inches(1.8), Inches(1.3)]
    col_x = [Inches(0.3)]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)

    add_rect(slide, Inches(0.3), Inches(2.0), sum(col_w), Inches(0.4), fill=ACCENT)
    for hdr, cx, cw in zip(headers, col_x, col_w):
        add_text(
            slide,
            hdr,
            cx + Inches(0.05),
            Inches(2.0),
            cw,
            Inches(0.4),
            size=Pt(9),
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER
        )

    for i, comp in enumerate(companies[:5]):
        row_y = Inches(2.4) + i * Inches(0.65)
        row_fill = BG_CARD if i % 2 == 0 else RGBColor(0x13, 0x18, 0x28)
        add_rect(slide, Inches(0.3), row_y, sum(col_w), Inches(0.6), fill=row_fill)

        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]
        add_rect(slide, Inches(0.3), row_y, Inches(0.05), Inches(0.6), fill=color)

        row_vals = [
            comp.get("company", "")[:20],
            comp.get("channel_name", "")[:22],
            fmt_num(comp.get("subscribers", 0)),
            fmt_num(comp.get("total_videos", 0)),
            fmt_num(comp.get("total_views", 0)),
            safe_str(comp.get("channel_since", "N/A"))[:7]
        ]

        for j, (val, cx, cw) in enumerate(zip(row_vals, col_x, col_w)):
            add_text(
                slide,
                val,
                cx + Inches(0.1),
                row_y + Inches(0.15),
                cw - Inches(0.1),
                Inches(0.35),
                size=Pt(10),
                color=color if j == 0 else WHITE,
                bold=(j == 0)
            )

    chart_data = ChartData()
    chart_data.categories = [c.get("company", "")[:15] for c in companies[:5]]
    chart_data.add_series("Subscribers", [c.get("subscribers", 0) for c in companies[:5]])

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.3),
        Inches(5.5),
        Inches(12.7),
        Inches(1.7),
        chart_data
    ).chart

    chart.has_legend = False
    chart.has_title = False

    for i, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = COMPANY_COLORS[i % len(COMPANY_COLORS)]


def slide_content_performance(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Content Performance", ACCENT3)
    slide_title(slide, "Top Performing Videos by Views & Engagement")
    add_accent_bar(slide, ACCENT3)

    companies = [c for c in report_data.get("companies", []) if "error" not in c]
    if not companies:
        add_text(slide, "No company data available.", Inches(0.5), Inches(2.0), Inches(12), Inches(0.5), size=Pt(14), color=WHITE)
        return

    col_w = Inches(13.0 / max(len(companies), 1))

    for i, comp in enumerate(companies[:4]):
        cx = Inches(0.2) + i * col_w
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]

        add_rect(slide, cx, Inches(1.9), col_w - Inches(0.15), Inches(0.35), fill=color)
        add_text(
            slide,
            comp.get("company", "")[:18],
            cx + Inches(0.1),
            Inches(1.92),
            col_w - Inches(0.25),
            Inches(0.3),
            size=Pt(10),
            bold=True,
            color=WHITE
        )

        top5 = comp.get("top_videos", [])[:4]
        for j, vid in enumerate(top5):
            vy = Inches(2.3) + j * Inches(1.15)
            add_rect(
                slide,
                cx,
                vy,
                col_w - Inches(0.15),
                Inches(1.1),
                fill=BG_CARD,
                line_color=color,
                line_width=Pt(0.5)
            )
            title = safe_str(vid.get("title", ""))
            title = title[:55] + ("…" if len(title) > 55 else "")
            add_text(
                slide,
                title,
                cx + Inches(0.08),
                vy + Inches(0.05),
                col_w - Inches(0.2),
                Inches(0.5),
                size=Pt(8.5),
                color=WHITE,
                wrap=True
            )
            add_text(
                slide,
                f"👁 {fmt_num(vid.get('views', 0))}   👍 {fmt_num(vid.get('likes', 0))}   💬 {fmt_num(vid.get('comments', 0))}",
                cx + Inches(0.08),
                vy + Inches(0.65),
                col_w - Inches(0.2),
                Inches(0.25),
                size=Pt(8),
                color=color
            )
            add_text(
                slide,
                f"Engagement: {vid.get('engagement_rate', 0)}%",
                cx + Inches(0.08),
                vy + Inches(0.85),
                col_w - Inches(0.2),
                Inches(0.2),
                size=Pt(8),
                color=GRAY
            )


def slide_topics_themes(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Content Topics & Themes", ACCENT4)
    slide_title(slide, "What Each Company Covers")
    add_accent_bar(slide, ACCENT4)

    companies = [c for c in report_data.get("companies", []) if "error" not in c]
    insights = report_data.get("insights", {})
    themes_map = insights.get("content_themes", {})

    col_w = Inches(12.7 / max(len(companies), 1))

    for i, comp in enumerate(companies[:5]):
        cx = Inches(0.3) + i * col_w
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]

        add_rect(slide, cx, Inches(1.9), col_w - Inches(0.1), Inches(0.35), fill=color)
        add_text(
            slide,
            comp.get("company", "")[:18],
            cx + Inches(0.08),
            Inches(1.92),
            col_w - Inches(0.2),
            Inches(0.3),
            size=Pt(10),
            bold=True,
            color=WHITE
        )

        topics = themes_map.get(comp.get("company", ""), []) or comp.get("top_topics", [])[:8]

        for j, topic in enumerate(topics[:8]):
            ty = Inches(2.4) + j * Inches(0.53)
            add_rect(slide, cx, ty, col_w - Inches(0.1), Inches(0.45), fill=BG_CARD)
            add_rect(slide, cx, ty, Inches(0.04), Inches(0.45), fill=color)
            add_text(
                slide,
                safe_str(topic)[:35],
                cx + Inches(0.12),
                ty + Inches(0.08),
                col_w - Inches(0.2),
                Inches(0.3),
                size=Pt(9),
                color=WHITE
            )


def slide_posting_frequency(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Posting Frequency & Consistency", ACCENT2)
    slide_title(slide, "Who Is Most Active and On What Cadence?")
    add_accent_bar(slide, ACCENT2)

    companies = [c for c in report_data.get("companies", []) if "error" not in c]
    insights = report_data.get("insights", {})

    for i, comp in enumerate(companies[:5]):
        cx = Inches(0.3) + i * Inches(2.55)
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]

        metric_card(
            slide,
            cx,
            Inches(1.9),
            Inches(2.4),
            Inches(1.1),
            "Company",
            comp.get("company", "")[:18],
            color=color
        )
        metric_card(
            slide,
            cx,
            Inches(3.1),
            Inches(2.4),
            Inches(1.0),
            "Posting Frequency",
            comp.get("posting_frequency", "N/A"),
            sub=f"{comp.get('total_videos', 0)} total videos",
            color=color
        )
        metric_card(
            slide,
            cx,
            Inches(4.2),
            Inches(2.4),
            Inches(0.9),
            "Channel Age",
            safe_str(comp.get("channel_since", "N/A"))[:7],
            color=color
        )

    insight_text = insights.get("posting_insight", "")
    if insight_text:
        add_rect(
            slide,
            Inches(0.3),
            Inches(5.3),
            Inches(12.7),
            Inches(1.5),
            fill=BG_CARD,
            line_color=ACCENT2,
            line_width=Pt(1)
        )
        add_text(
            slide,
            "💡 Insight",
            Inches(0.5),
            Inches(5.38),
            Inches(4),
            Inches(0.35),
            size=Pt(10),
            bold=True,
            color=ACCENT2
        )
        add_text(
            slide,
            insight_text,
            Inches(0.5),
            Inches(5.65),
            Inches(12.2),
            Inches(1.0),
            size=Pt(11),
            color=WHITE,
            wrap=True
        )


def slide_engagement_analysis(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Engagement Analysis", ACCENT)
    slide_title(slide, "Average Views, Likes & Comments Per Video")
    add_accent_bar(slide)

    companies = [c for c in report_data.get("companies", []) if "error" not in c]

    for i, comp in enumerate(companies[:5]):
        cx = Inches(0.3) + i * Inches(2.55)
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]

        add_rect(
            slide,
            cx,
            Inches(1.9),
            Inches(2.4),
            Inches(3.0),
            fill=BG_CARD,
            line_color=color,
            line_width=Pt(1)
        )
        add_text(
            slide,
            comp.get("company", "")[:18],
            cx + Inches(0.1),
            Inches(2.0),
            Inches(2.2),
            Inches(0.35),
            size=Pt(10),
            bold=True,
            color=color
        )

        metrics = [
            ("Avg Views", fmt_num(comp.get("avg_views", 0))),
            ("Avg Likes", fmt_num(comp.get("avg_likes", 0))),
            ("Avg Comments", fmt_num(comp.get("avg_comments", 0))),
            ("Engagement Rate", f"{comp.get('avg_engagement', 0)}%"),
        ]

        for j, (label, val) in enumerate(metrics):
            my = Inches(2.4) + j * Inches(0.6)
            add_text(
                slide,
                label,
                cx + Inches(0.1),
                my,
                Inches(2.2),
                Inches(0.25),
                size=Pt(8),
                color=GRAY
            )
            add_text(
                slide,
                val,
                cx + Inches(0.1),
                my + Inches(0.22),
                Inches(2.2),
                Inches(0.32),
                size=Pt(15),
                bold=True,
                color=WHITE
            )

    chart_data = ChartData()
    chart_data.categories = [c.get("company", "")[:12] for c in companies[:5]]
    chart_data.add_series("Avg Views", [c.get("avg_views", 0) for c in companies[:5]])
    chart_data.add_series("Avg Likes", [c.get("avg_likes", 0) for c in companies[:5]])

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.3),
        Inches(5.1),
        Inches(12.7),
        Inches(2.1),
        chart_data
    ).chart

    chart.has_title = False

    for i, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = [ACCENT, ACCENT2][i % 2]

    insights = report_data.get("insights", {})
    eng_insight = insights.get("engagement_insight", "")
    if eng_insight:
        add_text(
            slide,
            f"💡 {eng_insight}",
            Inches(0.3),
            Inches(7.0),
            Inches(12.7),
            Inches(0.35),
            size=Pt(9),
            color=GRAY,
            wrap=True
        )


def slide_gap_analysis(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Gap Analysis", ACCENT3)
    slide_title(slide, "Topics & Formats Competitors Are Missing")
    add_accent_bar(slide, ACCENT3)

    insights = report_data.get("insights", {})
    gaps = insights.get("content_gaps", [])
    missing_formats = insights.get("missing_formats", [])

    add_text(
        slide,
        "Content Gaps",
        Inches(0.4),
        Inches(1.9),
        Inches(6),
        Inches(0.4),
        size=Pt(13),
        bold=True,
        color=ACCENT3
    )

    for i, gap in enumerate(gaps[:6]):
        gy = Inches(2.35) + i * Inches(0.7)
        add_rect(
            slide,
            Inches(0.4),
            gy,
            Inches(5.8),
            Inches(0.6),
            fill=BG_CARD,
            line_color=ACCENT3,
            line_width=Pt(0.5)
        )
        add_rect(slide, Inches(0.4), gy, Inches(0.05), Inches(0.6), fill=ACCENT3)
        add_text(
            slide,
            safe_str(gap),
            Inches(0.6),
            gy + Inches(0.13),
            Inches(5.5),
            Inches(0.35),
            size=Pt(11),
            color=WHITE
        )

    add_text(
        slide,
        "Missing Formats",
        Inches(7.0),
        Inches(1.9),
        Inches(6),
        Inches(0.4),
        size=Pt(13),
        bold=True,
        color=ACCENT4
    )

    for i, fmt in enumerate(missing_formats[:6]):
        fy = Inches(2.35) + i * Inches(0.7)
        add_rect(
            slide,
            Inches(7.0),
            fy,
            Inches(5.8),
            Inches(0.6),
            fill=BG_CARD,
            line_color=ACCENT4,
            line_width=Pt(0.5)
        )
        add_rect(slide, Inches(7.0), fy, Inches(0.05), Inches(0.6), fill=ACCENT4)
        add_text(
            slide,
            safe_str(fmt),
            Inches(7.2),
            fy + Inches(0.13),
            Inches(5.5),
            Inches(0.35),
            size=Pt(11),
            color=WHITE
        )


def slide_recommendations(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Recommendations", ACCENT2)
    slide_title(slide, "Actionable Video Marketing Steps")
    add_accent_bar(slide, ACCENT2)

    insights = report_data.get("insights", {})
    recs = insights.get("recommendations", [])

    for i, rec in enumerate(recs[:5]):
        row = i // 3
        col = i % 3
        rx = Inches(0.3) + col * Inches(4.35)
        ry = Inches(1.9) + row * Inches(2.4)
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]

        add_rect(
            slide,
            rx,
            ry,
            Inches(4.1),
            Inches(2.2),
            fill=BG_CARD,
            line_color=color,
            line_width=Pt(1)
        )
        add_rect(slide, rx, ry, Inches(4.1), Inches(0.4), fill=color)

        num = f"0{i+1}"
        add_text(
            slide,
            num,
            rx + Inches(0.12),
            ry + Inches(0.05),
            Inches(0.5),
            Inches(0.32),
            size=Pt(14),
            bold=True,
            color=WHITE
        )

        title = rec.get("title", f"Recommendation {i+1}")
        detail = rec.get("detail", "")

        add_text(
            slide,
            safe_str(title)[:40],
            rx + Inches(0.12),
            ry + Inches(0.5),
            Inches(3.8),
            Inches(0.45),
            size=Pt(12),
            bold=True,
            color=WHITE
        )
        add_text(
            slide,
            safe_str(detail)[:200],
            rx + Inches(0.12),
            ry + Inches(0.95),
            Inches(3.8),
            Inches(1.15),
            size=Pt(9.5),
            color=GRAY,
            wrap=True
        )


def slide_summary_scores(prs, report_data):
    slide = add_slide(prs)
    section_header(slide, "Summary & Rankings", ACCENT)
    slide_title(slide, "Company Scorecard — Key Metrics Ranked")
    add_accent_bar(slide)

    insights = report_data.get("insights", {})
    scores = insights.get("company_scores", {})
    rankings = insights.get("rankings", [])
    companies = [c for c in report_data.get("companies", []) if "error" not in c]

    headers = ["Rank", "Company", "Content Quality", "Consistency", "Engagement", "Growth Potential", "Overall"]
    col_w = [Inches(0.7), Inches(2.0), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.9), Inches(1.6)]
    col_x = [Inches(0.3)]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)

    add_rect(slide, Inches(0.3), Inches(2.0), sum(col_w), Inches(0.4), fill=ACCENT)
    for hdr, cx, cw in zip(headers, col_x, col_w):
        add_text(
            slide,
            hdr,
            cx + Inches(0.05),
            Inches(2.0),
            cw,
            Inches(0.4),
            size=Pt(8.5),
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER
        )

    sorted_companies = companies[:]
    if rankings:
        rank_map = {name: i for i, name in enumerate(rankings)}
        sorted_companies.sort(key=lambda c: rank_map.get(c.get("company", ""), 99))

    for i, comp in enumerate(sorted_companies[:5]):
        row_y = Inches(2.4) + i * Inches(0.65)
        color = COMPANY_COLORS[i % len(COMPANY_COLORS)]
        row_fill = BG_CARD if i % 2 == 0 else RGBColor(0x13, 0x18, 0x28)

        add_rect(slide, Inches(0.3), row_y, sum(col_w), Inches(0.6), fill=row_fill)
        add_rect(slide, Inches(0.3), row_y, Inches(0.05), Inches(0.6), fill=color)

        comp_scores = scores.get(comp.get("company", ""), {})
        rank_str = f"#{i+1}"

        row_vals = [
            rank_str,
            comp.get("company", "")[:18],
            f"{comp_scores.get('content_quality', '—')}/10",
            f"{comp_scores.get('consistency', '—')}/10",
            f"{comp_scores.get('engagement', '—')}/10",
            f"{comp_scores.get('growth_potential', '—')}/10",
            f"{comp_scores.get('overall', '—')}/10",
        ]

        for j, (val, cx, cw) in enumerate(zip(row_vals, col_x, col_w)):
            is_overall = j == len(row_vals) - 1
            add_text(
                slide,
                val,
                cx + Inches(0.08),
                row_y + Inches(0.15),
                cw - Inches(0.1),
                Inches(0.35),
                size=Pt(10 if not is_overall else 12),
                bold=(j <= 1 or is_overall),
                color=color if j == 1 else (ACCENT2 if is_overall else WHITE),
                align=PP_ALIGN.CENTER
            )

    add_rect(
        slide,
        Inches(0.3),
        Inches(6.0),
        Inches(12.7),
        Inches(1.1),
        fill=BG_CARD,
        line_color=ACCENT,
        line_width=Pt(1)
    )
    add_text(
        slide,
        "Final Rankings",
        Inches(0.5),
        Inches(6.08),
        Inches(3),
        Inches(0.35),
        size=Pt(10),
        bold=True,
        color=ACCENT
    )
    ranking_text = "   ".join([f"#{j+1} {name}" for j, name in enumerate(rankings[:5])])
    add_text(
        slide,
        ranking_text,
        Inches(0.5),
        Inches(6.4),
        Inches(12.2),
        Inches(0.55),
        size=Pt(13),
        bold=True,
        color=WHITE
    )


# ── Master builder ───────────────────────────────────────────────

def build_pptx(report_data, output_path):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    slide_cover(prs, report_data)
    slide_executive_summary(prs, report_data)
    slide_channel_overview(prs, report_data)
    slide_content_performance(prs, report_data)
    slide_topics_themes(prs, report_data)
    slide_posting_frequency(prs, report_data)
    slide_engagement_analysis(prs, report_data)
    slide_gap_analysis(prs, report_data)
    slide_recommendations(prs, report_data)
    slide_summary_scores(prs, report_data)

    prs.save(output_path)
    return output_path